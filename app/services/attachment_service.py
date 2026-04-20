"""Extract text from email attachments (PDF, Excel, Word)."""

from __future__ import annotations

import io
import logging
import re
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_text_from_bytes(data: bytes, filename: str) -> str:
    """Dispatch to the right extractor based on file extension."""
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".pdf":
            return _extract_pdf(data)
        if ext in (".xlsx", ".xls"):
            return _extract_excel(data)
        if ext in (".docx", ".doc"):
            return _extract_docx(data)
        if ext in (".csv", ".txt", ".log"):
            return data.decode("utf-8", errors="replace")
    except Exception as exc:
        logger.warning("Failed to extract text from %s: %s", filename, exc)
    return ""


def extract_text_from_file(path: str) -> str:
    """Read a local file and extract text."""
    p = Path(path)
    if not p.exists():
        return ""
    return extract_text_from_bytes(p.read_bytes(), p.name)


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        logger.warning("pypdf not installed — cannot read PDFs")
        return ""
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


def _extract_excel(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        logger.warning("openpyxl not installed — cannot read Excel files")
        return ""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        lines.append(f"--- Sheet: {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                lines.append("\t".join(cells))
    wb.close()
    return "\n".join(lines)


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError:
        logger.warning("python-docx not installed — cannot read Word files")
        return ""
    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


# ── Price extraction helpers ──────────────────────────────────────────────────

_PRICE_PATTERNS = [
    # Rs / INR patterns: Rs 150, Rs. 150.00, INR 1,50,000, Rs 5,00,000/-
    re.compile(
        r"(?:Rs\.?|INR|₹)\s*([\d,]+(?:\.\d{1,2})?)\s*(?:/[-–]|per\s+\w+)?",
        re.IGNORECASE,
    ),
    # "price" or "rate" followed by number
    re.compile(
        r"(?:price|rate|cost|amount|value)\s*[:=]?\s*(?:Rs\.?|INR|₹)?\s*([\d,]+(?:\.\d{1,2})?)",
        re.IGNORECASE,
    ),
]


def extract_prices(text: str) -> list[dict]:
    """Pull price mentions from text. Returns [{value, currency, context}]."""
    results: list[dict] = []
    seen: set[str] = set()
    for pat in _PRICE_PATTERNS:
        for m in pat.finditer(text):
            raw = m.group(1).replace(",", "")
            try:
                val = float(raw)
            except ValueError:
                continue
            if val < 1 or raw in seen:
                continue
            seen.add(raw)
            start = max(0, m.start() - 40)
            end = min(len(text), m.end() + 40)
            ctx = text[start:end].replace("\n", " ").strip()
            results.append({"value": val, "currency": "INR", "context": ctx})
    return results


def extract_product_mentions(text: str) -> list[str]:
    """Pull potential product/material names from text (heuristic)."""
    products: list[str] = []
    patterns = [
        re.compile(
            r"(?:product|material|item|grade|type)\s*[:=]?\s*([A-Za-z0-9][\w\s\-/]{2,30})",
            re.IGNORECASE,
        ),
        re.compile(
            r"(?:for|of|regarding)\s+([A-Z][\w\s\-/]{2,25}?)(?:\s+(?:at|@|price|rate|order))",
            re.IGNORECASE,
        ),
    ]
    for pat in patterns:
        for m in pat.finditer(text):
            name = m.group(1).strip()
            if name and len(name) > 2 and name not in products:
                products.append(name)
    return products[:10]
